import calendar
import io
import json
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.styles.fonts import DEFAULT_FONT

try:
  import holidays
except ImportError:
  holidays = None

st.set_page_config(
    page_title="ระบบสรุปวันลาประจำเดือน", page_icon="📅", layout="wide"
)
st.title("📌 ระบบสรุปวันลาประจำเดือน (Template ราชการ)")

# แสดงข้อความแจ้งผลที่ฝากไว้ก่อน rerun
if msg := st.session_state.pop("flash", None):
  st.toast(msg)

# -------------------------------------------------------------------------
# 1. ค่าคงที่
# -------------------------------------------------------------------------
SEED_EMPLOYEES = {
    "พนักงานจ้างเหมาเอกชน": [
        {"name": "นางสาวณัฐพร  มิตรพิทักษ์", "position": "พนักงานจ้างปฏิบัติงานนโยบายและแผน"},
        {"name": "นางสาวณัฐธิชา  แถวทัด", "position": "พนักงานจ้างปฏิบัติงานนิเทศติดตาม\nและประเมินผล"},
        {"name": "นางสาววิชุดา  ศรีคร้ามครัน", "position": "พนักงานจ้างปฏิบัติงานธุรการ (กลุ่มพัฒนาฯ)"},
        {"name": "นางสาวยัสมี  ดาโอะ", "position": "พนักงานจ้างปฏิบัติงานธุรการ"},
        {"name": "นางสาวภาวินีย์  ตันสกุล", "position": "พนักงานจ้างทำความสะอาด"},
        {"name": "นายภานุวัฒน์  จันทร์ชีวา", "position": "พนักงานจ้างปฏิบัติงานขับรถยนต์"},
        {"name": "นางสาวหัทยาภักดิ์ แสงแก้ว", "position": "พนักงานจ้างปฏิบัติงานการเงินและบัญชี"},
        {"name": "นายหนึ่ง  สุนทรรัตน์", "position": "พนักงานจ้างปฏิบัติงานรักษาความปลอดภัย"},
    ],
    "พนักงานราชการ": [
        {"name": "นางสาวอมรารัตน์  พูลสวัสดิ์", "position": "พนักงานจัดการงานบุคคล"},
        {"name": "นายสมชาย  สุพรดวงชัย", "position": "พนักงานจัดการงานทั่วไป"},
    ],
}

THAI_MONTHS = [
    "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม",
]
COLUMNS = [
    "ที่", "ชื่อ - สกุล", "ตำแหน่ง", "ตั้งแต่วันที่", "ถึงวันที่",
    "ลาป่วย", "ลากิจ", "รวม (วัน)", "หมายเหตุ",
]
SAVE_FILE = Path(__file__).resolve().parent / "flowleave_data.json"


# -------------------------------------------------------------------------
# 2. ฟังก์ชันช่วย
# -------------------------------------------------------------------------
def format_thai_date(d):
  return f"{d.day} {THAI_MONTHS[d.month - 1]} {d.year + 543}"


def fmt_days(v):
  return f"{v:g}" if v else "-"  # 1.0 -> "1", 4.5 -> "4.5", 0 -> "-"


def load_data():
  """ไฟล์ JSON คือแหล่งข้อมูลหลักเพียงแหล่งเดียว"""
  try:
    raw = json.loads(SAVE_FILE.read_text(encoding="utf-8"))
    records = {
        name: [
            {
                "note": "-",
                **r,
                "start": date.fromisoformat(r["start"]),
                "end": date.fromisoformat(r["end"]),
            }
            for r in recs
        ]
        for name, recs in raw.get("records", {}).items()
    }
    return raw.get("master") or SEED_EMPLOYEES, records
  except (FileNotFoundError, json.JSONDecodeError, KeyError, ValueError):
    return SEED_EMPLOYEES, {}


def save_data():
  payload = {
      "master": employees,
      "records": {
          name: [
              {**r, "start": r["start"].isoformat(), "end": r["end"].isoformat()}
              for r in recs
          ]
          for name, recs in records.items()
          if recs
      },
  }
  try:
    SAVE_FILE.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
  except OSError:
    st.session_state.flash = "⚠️ บันทึกไฟล์ข้อมูลไม่ได้ — ตรวจสอบสิทธิ์การเขียนโฟลเดอร์"


def rerun_with(msg, reopen=None):
  """ฝากข้อความแจ้งผล (และชื่อที่ต้องเปิดป๊อปอัปซ้ำ) แล้ว rerun"""
  st.session_state.setdefault("flash", msg)  # ไม่ทับข้อความ error จาก save_data
  st.session_state.reopen = reopen
  st.rerun()


def calc_leave_days(start, end, note):
  """นับวันทำการ (หักเสาร์-อาทิตย์และวันหยุดราชการ) และหัก 0.5 ถ้าระบุครึ่งวัน"""
  th_holidays = holidays.TH(years=start.year) if holidays else set()
  dates = (start + timedelta(days=i) for i in range((end - start).days + 1))
  days = max(sum(1 for d in dates if d.weekday() < 5 and d not in th_holidays), 1)
  if "ครึ่งวัน" in note or "0.5" in note:
    days -= 0.5
  return days


def month_records(name):
  """เฉพาะรายการลาที่อยู่ในเดือนที่เลือก"""
  return [r for r in records.get(name, []) if month_start <= r["start"] <= month_end]


# -------------------------------------------------------------------------
# 3. โหลดข้อมูล + ส่วนหัว
# -------------------------------------------------------------------------
employees, records = load_data()

with st.sidebar:
  st.header("⚙️ คำแนะนำการใช้งาน")
  st.markdown(
      "1. **เลือกงวดเวลา:** กำหนดเดือนและปีที่ต้องการสรุปรายงาน\n"
      "2. **บันทึกวันลา:** แตะแถวชื่อพนักงานในตารางเพื่อเปิดหน้าต่างจัดการ\n"
      "3. **ครึ่งวัน:** พิมพ์คำว่า 'ครึ่งวัน' ในหมายเหตุ ระบบหัก 0.5 วันให้\n"
      "4. **ดาวน์โหลด:** ตรวจสอบตารางแล้วกดดาวน์โหลดไฟล์ Excel"
  )

today = date.today()
st.markdown("### 🗓️ กำหนดงวดรายงานประจำเดือน")
c1, c2, c3 = st.columns([2, 2, 3])
selected_month = c1.selectbox("ประจำเดือน", THAI_MONTHS, index=today.month - 1)
selected_year = int(c2.number_input("ปี พ.ศ.", value=today.year + 543, step=1))
employee_type = c3.selectbox("ประเภทพนักงาน", list(employees))

month_num = THAI_MONTHS.index(selected_month) + 1
ce_year = selected_year - 543
month_start = date(ce_year, month_num, 1)
month_end = date(ce_year, month_num, calendar.monthrange(ce_year, month_num)[1])

st.info(
    f"📂 กำลังจัดการข้อมูลของ: **{employee_type}** "
    f"ประจำเดือน**{selected_month} {selected_year}**"
)


# -------------------------------------------------------------------------
# 4. ตารางสรุป
# -------------------------------------------------------------------------
def build_summary_dataframe():
  rows = []
  for idx, emp in enumerate(employees.get(employee_type, []), 1):
    recs = month_records(emp["name"])
    sick = sum(r["days"] for r in recs if r["category"] == "ลาป่วย")
    kit = sum(r["days"] for r in recs if r["category"] == "ลากิจ")
    rows.append({
        "ที่": idx,
        "ชื่อ - สกุล": emp["name"],
        "ตำแหน่ง": emp["position"],
        "ตั้งแต่วันที่": "\n".join(format_thai_date(r["start"]) for r in recs) or "-",
        "ถึงวันที่": "\n".join(format_thai_date(r["end"]) for r in recs) or "-",
        "ลาป่วย": fmt_days(sick),
        "ลากิจ": fmt_days(kit),
        "รวม (วัน)": fmt_days(sick + kit),
        "หมายเหตุ": "\n".join(
            f"{r['note'] or '-'} ({r['days']:g} วัน)" for r in recs
        ) or "-",
    })
  return pd.DataFrame(rows, columns=COLUMNS)


# -------------------------------------------------------------------------
# 5. ป๊อปอัปจัดการวันลา
# -------------------------------------------------------------------------
@st.dialog("📅 จัดการประวัติการลาบุคลากร", width="large")
def open_leave_dialog(emp_name):
  st.markdown(f"### 👤 {emp_name}")
  st.markdown(f"#### 📋 รายการลาเดือน{selected_month} {selected_year}")

  recs = month_records(emp_name)
  if not recs:
    st.info("ยังไม่มีประวัติการลาในเดือนนี้")
  for i, rec in enumerate(recs):
    a, b, c = st.columns([3, 6, 1])
    a.markdown(f"**{rec['category']}** ({rec['days']:g} วัน)")
    b.text(
        f"{format_thai_date(rec['start'])} ถึง "
        f"{format_thai_date(rec['end'])} [{rec['note']}]"
    )
    if c.button("🗑️", key=f"del_{emp_name}_{i}", help="ลบรายการนี้"):
      records[emp_name].remove(rec)
      save_data()
      rerun_with("ลบรายการสำเร็จ", reopen=emp_name)

  st.divider()
  st.markdown("#### ➕ เพิ่มรายการลาใหม่")
  category = st.selectbox("ประเภทการลา", ["ลาป่วย", "ลากิจ"], key="d_cat")
  a, b = st.columns(2)
  d_start = a.date_input(
      "ตั้งแต่วันที่", value=month_start,
      min_value=month_start, max_value=month_end, key="d_start",
  )
  d_end = b.date_input(
      "ถึงวันที่", value=month_start,
      min_value=month_start, max_value=month_end, key="d_end",
  )
  note = st.text_input(
      "หมายเหตุ / ลักษณะการลา", value="-",
      placeholder="เช่น ลาครึ่งวันเช้า, วันแรกเต็ม วันที่สองครึ่งวันบ่าย",
      key="d_note",
  ).strip() or "-"

  if d_end < d_start:
    st.error("⚠️ วันที่สิ้นสุดต้องไม่ก่อนวันที่เริ่มต้น")
    return

  days = calc_leave_days(d_start, d_end, note)
  st.info(f"📌 **จำนวนวันที่คำนวณได้:** `{days:g} วัน`")

  if st.button("➕ บันทึกเพิ่มรายการนี้", type="primary", use_container_width=True):
    records.setdefault(emp_name, []).append({
        "start": d_start, "end": d_end,
        "category": category, "days": days, "note": note,
    })
    save_data()
    rerun_with("เพิ่มรายการลาสำเร็จ", reopen=emp_name)


# -------------------------------------------------------------------------
# 6. แสดงตาราง + เปิดป๊อปอัปเฉพาะเมื่อเลือกแถวใหม่
# -------------------------------------------------------------------------
st.subheader("📊 ตารางสรุปข้อมูลบุคลากรและวันลาประจำเดือน")
st.markdown("👉 **แตะแถวพนักงานในตาราง เพื่อเปิดหน้าต่างจัดการประวัติการลา**")

df_display = build_summary_dataframe()
if df_display.empty:
  st.info("📭 ยังไม่มีบุคลากรในกลุ่มนี้ — เพิ่มได้จากเมนูด้านล่าง")

event = st.dataframe(
    df_display,
    use_container_width=True,
    hide_index=True,
    selection_mode="single-row",
    on_select="rerun",
    key="employee_table",
)

sel = event.selection.rows
picked = (
    df_display.iloc[sel[0]]["ชื่อ - สกุล"]
    if sel and sel[0] < len(df_display) else None
)
target = st.session_state.pop("reopen", None) or (
    picked if picked != st.session_state.get("last_picked") else None
)
st.session_state.last_picked = picked
if target:
  open_leave_dialog(target)

# -------------------------------------------------------------------------
# 7. เพิ่ม / ลบรายชื่อพนักงาน
# -------------------------------------------------------------------------
with st.expander("➕ เพิ่มรายชื่อพนักงานใหม่ / 🗑️ ลบรายชื่อพนักงาน"):
  add_col, del_col = st.columns(2)

  with add_col:
    with st.form("add_form", clear_on_submit=True):
      new_n = st.text_input("ชื่อ - สกุลใหม่").strip()
      new_p = st.text_input("ตำแหน่งใหม่").strip()
      if st.form_submit_button("➕ เพิ่มพนักงาน") and new_n:
        all_names = {e["name"] for group in employees.values() for e in group}
        if new_n in all_names:
          st.warning("มีชื่อนี้อยู่ในระบบแล้ว")
        else:
          employees.setdefault(employee_type, []).append(
              {"name": new_n, "position": new_p or "-"}
          )
          save_data()
          rerun_with("เพิ่มรายชื่อสำเร็จ")

  names = [e["name"] for e in employees.get(employee_type, [])]
  if names:
    with del_col:
      with st.form("del_form"):
        del_n = st.selectbox("เลือกรายชื่อที่ต้องการลบออก", names)
        if st.form_submit_button("🗑️ ลบพนักงานออก"):
          employees[employee_type] = [
              e for e in employees[employee_type] if e["name"] != del_n
          ]
          records.pop(del_n, None)
          save_data()
          rerun_with("ลบรายชื่อสำเร็จ")

# -------------------------------------------------------------------------
# 8. สร้างไฟล์ Excel ตามแบบฟอร์มราชการ
# -------------------------------------------------------------------------
DEFAULT_FONT.name = "TH SarabunPSK"
DEFAULT_FONT.size = 14


def generate_excel(df, month, year, emp_type):
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "สรุปวันลา"

  side = Side(style="thin", color="000000")
  border = Border(left=side, right=side, top=side, bottom=side)
  center = Alignment(horizontal="center", vertical="center", wrap_text=True)
  left = Alignment(horizontal="left", vertical="center", wrap_text=True)
  f_title = Font(name="TH SarabunPSK", size=16, bold=True)
  f_bold = Font(name="TH SarabunPSK", size=14, bold=True)
  f_normal = Font(name="TH SarabunPSK", size=14)

  for col, width in zip("ABCDEFGHI", [6, 28, 42, 20, 20, 12, 12, 12, 28]):
    ws.column_dimensions[col].width = width

  # หัวเรื่อง 2 บรรทัด
  titles = [
      f"ตารางสรุปวันลา {emp_type} สำนักงานศึกษาธิการจังหวัดสมุทรสงคราม",
      f"ประจำเดือน{month} {year}",
  ]
  for row, text in enumerate(titles, start=1):
    ws.merge_cells(f"A{row}:I{row}")
    cell = ws[f"A{row}"]
    cell.value, cell.font, cell.alignment = text, f_title, center

  # หัวตาราง 2 ชั้น
  header_cols = zip(
      "ABCDEHI",
      ["ที่", "ชื่อ - สกุล", "ตำแหน่ง", "ตั้งแต่วันที่", "ถึงวันที่", "รวม\n(วัน)", "หมายเหตุ"],
  )
  for col, text in header_cols:
    ws.merge_cells(f"{col}4:{col}5")
    ws[f"{col}4"] = text
  ws.merge_cells("F4:G4")
  ws["F4"], ws["F5"], ws["G5"] = "ครั้งนี้", "ลาป่วย", "ลากิจ"

  for row in ws.iter_rows(min_row=4, max_row=5, max_col=9):
    for cell in row:
      cell.font, cell.alignment, cell.border = f_bold, center, border

  # ข้อมูล (ลำดับคอลัมน์ของ df ตรงกับ A–I)
  for r, values in enumerate(df.itertuples(index=False), start=6):
    for c, value in enumerate(values, start=1):
      cell = ws.cell(row=r, column=c, value=value)
      cell.font, cell.border = f_normal, border
      cell.alignment = left if c in (2, 3, 9) else center

  # ช่องลงชื่อ
  sig = ws.max_row + 3
  ws.merge_cells(f"F{sig}:I{sig}")
  cell = ws[f"F{sig}"]
  cell.value = "(ลงชื่อ)....................................................ผู้ตรวจสอบ"
  cell.font, cell.alignment = f_normal, center

  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue()


st.subheader("📥 ดาวน์โหลดรายงานประจำเดือน")
st.download_button(
    "📥 ดาวน์โหลดไฟล์ Excel ตามแบบฟอร์มราชการ",
    data=generate_excel(df_display, selected_month, selected_year, employee_type),
    file_name=f"สรุปวันลา_{employee_type}_{selected_month}_{selected_year}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    type="primary",
)